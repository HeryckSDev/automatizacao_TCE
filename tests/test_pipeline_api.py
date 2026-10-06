import unittest, tempfile, base64
from pathlib import Path
from io import BytesIO
from docx import Document
from openpyxl import load_workbook
from app import create_app

class PipelineApiTest(unittest.TestCase):
    def test_docx_componentes_gera_excel_e_rejeita_pspn_no_total(self):
        doc=Document();doc.add_paragraph('Município: Cidade Exemplo\nAno: 2026\nCargo: Cargo Exemplo\nJornada: 40 horas')
        table=doc.add_table(rows=3,cols=4)
        for row,values in zip(table.rows,[['Classe','Nível','VB','Total'],['A','1','2000,00','2400,00'],['A','2','2100,00','2520,00']]):
            for cell,value in zip(row.cells,values):cell.text=value
        b=BytesIO();doc.save(b)
        with tempfile.TemporaryDirectory() as tmp:
            app=create_app(None,tmp)
            app.config['TESTING']=True;client=app.test_client()
            r=client.post('/api/automacao/preparar-tabela',data={'arquivo':(BytesIO(b.getvalue()),'teste.docx')})
            self.assertEqual(r.status_code,200,r.json)
            # Component handling applies to native PDF tables; DOCX must use the same component reader.
            self.assertEqual(len(r.json['candidatos']),2)
            pedido={'preparacao_id':r.json['preparacao_id'],'candidato':1,'confirmar_recorte':True,'complementos':{'municipio':'Cidade Exemplo','profissao':'Cargo Exemplo','ano':2026,'jornada_semanal':40},'comparacao':{'tipo':'percentual','valor':5}}
            export=client.post('/api/automacao/salarios',json=pedido)
            self.assertEqual(export.status_code,200,export.json)
            wb=load_workbook(BytesIO(base64.b64decode(export.json['arquivo']['conteudo_base64'])),data_only=True)
            fontes=[v for row in wb['Fontes e premissas'].values for v in row if isinstance(v,str)]
            self.assertIn('SHA-256 do documento',fontes)
            pedido['comparacao']={'tipo':'pspn'}
            self.assertEqual(client.post('/api/automacao/salarios',json=pedido).status_code,400)

    def test_bloqueio_pspn_tambem_na_camada_de_calculo(self):
        import json
        from src.salarial.analise_service import analisar_tabela_salarial
        fonte=Path(__file__).resolve().parents[1]/'data/processed/cafezal_do_sul_2026.json'
        dados=json.loads(fonte.read_text());dados['pspn_inaplicavel']=True
        with self.assertRaisesRegex(ValueError,'referência própria'):
            analisar_tabela_salarial(dados)
