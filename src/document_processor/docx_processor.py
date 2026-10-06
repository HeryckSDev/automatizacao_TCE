from pathlib import Path
from io import BytesIO
import shutil,subprocess,tempfile
from docx import Document
from src.salarial.leitor_geral import ler_matrizes,ler_texto
from src.salarial.documento_service import _texto_doc_legado
from .table_processor import ler_tabelas

def ler_word(nome,conteudo,organizar):
    ext=Path(nome).suffix.lower()
    if ext=='.doc':
        exe=shutil.which('soffice') or shutil.which('libreoffice')
        if exe:
            with tempfile.TemporaryDirectory(prefix='jornada-word-') as tmp:
                pasta=Path(tmp);src=pasta/'original.doc';src.write_bytes(conteudo)
                proc=subprocess.run([exe,f'-env:UserInstallation={(pasta / "perfil").as_uri()}', '--headless','--convert-to','docx','--outdir',str(pasta),str(src)],capture_output=True,timeout=45)
                convertido=pasta/'original.docx'
                if proc.returncode==0 and convertido.exists():
                    d,t,p=ler_word('convertido.docx',convertido.read_bytes(),organizar)
                    for c in d:c['arquivo_fonte']=c['arquivo_fonte'].replace('convertido.docx',nome);c['avisos_importacao'].append('DOC convertido localmente para DOCX; confira a conversão no original.')
                    return d,t,p
        text=_texto_doc_legado(conteudo)
        d,p=ler_texto(text,nome,organizar);return d,text,p
    try:doc=Document(BytesIO(conteudo))
    except Exception as erro:raise ValueError('DOCX inválido ou protegido; envie um documento válido.') from erro
    text='\n'.join(p.text for p in doc.paragraphs)
    tables=[[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]
    d,p=ler_tabelas(tables,nome,organizar,text)
    metodo='tabela_docx'
    if not d and not p:
        d,p=ler_texto(text,nome,organizar);metodo='texto_docx'
    for c in d:c['metodo_extracao']=metodo
    return d,text,p

