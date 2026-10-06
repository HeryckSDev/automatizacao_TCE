"""Leitura estruturada: não executa macros nem calcula fórmulas."""
from io import BytesIO
from openpyxl import load_workbook

def extrair_xlsx(conteudo):
    wb = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    try:
        tabelas = []
        for ws in wb:
            if ws.max_row>20000 or ws.max_column>220:
                raise ValueError("Planilha extensa: exporte somente a tabela selecionada.")
            tabelas.append(list(ws.values))
        return tabelas
    finally: wb.close()
